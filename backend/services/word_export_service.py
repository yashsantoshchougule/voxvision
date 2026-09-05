from io import BytesIO
from docx import Document
from docx.shared import Pt

def build_docx(meeting: dict, report: dict, important_points: list[dict], screen_items: list[dict]) -> BytesIO:
    document = Document()
    heading = document.add_heading("VoxVision AI Meeting Report", level=0)
    heading.runs[0].font.size = Pt(22)
    document.add_paragraph(f"Meeting title: {meeting.get('title', 'Untitled')}\nPlatform: {meeting.get('platform', 'Google Meet')}\nOutput mode: {meeting.get('mode') or meeting.get('selected_mode', 'combined')}\nStarted: {meeting.get('started_at', '')}\nEnded: {meeting.get('ended_at', '')}")
    for title, key in (("Summary", "summary"), ("Detailed notes", "notes"), ("Insights", "insights"), ("Decisions", "decisions"), ("Deadlines", "deadlines"), ("Graph insights", "graph_insights"), ("Recommendations", "recommendations"), ("Uncertain information", "uncertain_information")):
        document.add_heading(title, level=1)
        values = report.get(key) if key != "summary" else [report.get(key) or "Not specified"]
        for value in values or ["None recorded."]:
            document.add_paragraph(str(value), style="List Bullet" if key != "summary" else None)
    document.add_heading("Tasks", level=1)
    table = document.add_table(rows=1, cols=4)
    table.style = "Light Shading Accent 1"
    for cell, label in zip(table.rows[0].cells, ("Task", "Assigned To", "Deadline", "Priority")):
        cell.text = label
    for task in report.get("tasks", []):
        cells = table.add_row().cells
        for index, value in enumerate((task.get("task", ""), task.get("assigned_to") or "Not specified", task.get("deadline") or "Not specified", task.get("priority") or "Not specified")):
            cells[index].text = str(value)
    document.add_heading("Important points", level=1)
    for item in important_points or [{"label": "None", "transcript_context": "None recorded."}]:
        document.add_paragraph(f"{item.get('label')}: {item.get('user_note') or item.get('transcript_context') or item.get('screen_context')}", style="List Bullet")
    document.add_heading("Screen observations", level=1)
    for item in screen_items or [{"extracted_text": "None recorded."}]:
        document.add_paragraph(f"{item.get('extracted_text', '')} (type: {item.get('content_type', 'unknown')})", style="List Bullet")
    output = BytesIO()
    document.save(output)
    output.seek(0)
    return output
